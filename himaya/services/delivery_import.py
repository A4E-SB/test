"""
Bulk delivery-status import: load the CSV/Excel exported by Yalidine /
ZR Express / Maystro (or paste lines 'phone ; status') and update many
orders at once. Column names are auto-detected; status keywords are mapped
(fr/en/dz slang) to Himaya statuses.
"""

from __future__ import annotations

import csv
from pathlib import Path

from ..database.db import Database
from ..models import orders as orders_model
from ..services.phone import normalize_phone
from ..services import trust

# keyword -> Himaya status (checked lowercase, in order)
STATUS_MAP = [
    (("livre", "livré", "delivered", "deliver", "remis"), "delivered"),
    (("paye", "payé", "paid", "encaisse", "encaissé"), "paid"),
    (("refus", "refuse", "refused", "rejet"), "refused"),
    (("retourne", "retourné", "returned", "return", "introuvable", "pas repondu",
      "no answer", "pas de reponse", "hors ligne", "eteint", "éteint"), "ghosted"),
    (("annule", "annulé", "canceled", "cancelled"), "canceled"),
    (("expedie", "expédié", "shipped", "en cours", "transit", "en route",
      "pret", "prêt"), "shipped"),
    (("confirme", "confirmé", "confirmed"), "confirmed"),
]

PHONE_HEADERS = ("phone", "tel", "telephone", "téléphone", "gsm", "mobile", "numero", "numéro")
STATUS_HEADERS = ("statut", "status", "etat", "état", "situation", "commentaire", "remarque")


def _fold(s: str) -> str:
    """Lowercase + strip accents so 'Répondu' matches 'repondu'."""
    import unicodedata
    s = (s or "").lower()
    return "".join(c for c in unicodedata.normalize("NFD", s)
                   if unicodedata.category(c) != "Mn")


def map_status(raw: str) -> str | None:
    low = _fold(raw)
    if not low:
        return None
    for keywords, status in STATUS_MAP:
        if any(_fold(k) in low for k in keywords):
            return status
    return None


def _read_rows(path: Path) -> list[dict]:
    """Read a CSV or XLSX into normalized {header: value} dicts."""
    suffix = path.suffix.lower()
    if suffix in (".xlsx", ".xls"):
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        wb.close()
        if not rows:
            return []
        headers = [str(h or "").strip() for h in rows[0]]
        out = []
        for row in rows[1:]:
            rec = {}
            for i, h in enumerate(headers):
                if i < len(row):
                    rec[h] = "" if row[i] is None else str(row[i])
            out.append(rec)
        return out
    # CSV (excel export: often ';')
    for enc in ("utf-8-sig", "latin-1"):
        try:
            with open(path, encoding=enc, newline="") as f:
                sample = f.read(4096)
                f.seek(0)
                try:
                    dialect = csv.Sniffer().sniff(sample, delimiters=";,\t")
                except csv.Error:
                    dialect = csv.excel
                    dialect.delimiter = ";"
                return [dict(r) for r in csv.DictReader(f, dialect=dialect)]
        except UnicodeDecodeError:
            continue
    return []


def _find_col(row: dict, candidates) -> str | None:
    low = {k.lower(): k for k in row}
    for cand in candidates:
        if cand in low:
            return low[cand]
        # partial match too (e.g. "Téléphone client")
        for k_l, k in low.items():
            if cand in k_l:
                return k
    return None


def parse_delivery_file(path: str | Path) -> list[dict]:
    """
    Extract [{phone, status(himaya), raw_status, order_id?}] from a
    courier export. Also supports the app's own exported CSV (has 'status'
    and 'phone' columns) — in that case order ids come along.
    """
    rows = _read_rows(Path(path))
    out = []
    for row in rows:
        if not row:
            continue
        phone_col = _find_col(row, PHONE_HEADERS)
        status_col = _find_col(row, STATUS_HEADERS)
        if not phone_col:
            continue
        phone = normalize_phone(row[phone_col]) or ""
        raw = row.get(status_col, "") if status_col else ""
        status = map_status(raw)
        oid = None
        id_col = _find_col(row, ("id", "order", "commande", "reference", "référence"))
        if id_col and (row.get(id_col) or "").strip().isdigit():
            oid = int(row[id_col].strip())
        if phone or oid:
            out.append({"phone": phone, "status": status, "raw_status": raw,
                        "order_id": oid})
    return out


def parse_paste(text: str) -> list[dict]:
    """Lines like '0555123456 ; livré' or '0555123456 livré'."""
    out = []
    for line in (text or "").splitlines():
        line = line.strip()
        if not line:
            continue
        parts = [x.strip() for x in line.replace(";", " ").replace(",", " ").split(None, 1)]
        if not parts:
            continue
        phone = normalize_phone(parts[0]) or ""
        raw = parts[1] if len(parts) > 1 else ""
        out.append({"phone": phone, "status": map_status(raw), "raw_status": raw,
                    "order_id": None})
    return out


def apply_updates(db: Database, updates: list[dict]) -> dict:
    """
    Apply parsed updates: match each phone to its most recent open order
    (pending/confirmed/waiting_deposit/shipped) and set the new status.
    Returns {updated, unmatched, unmapped}.
    """
    updated = unmatched = unmapped = 0
    for u in updates:
        status = u.get("status")
        if not status:
            unmapped += 1
            continue
        oid = u.get("order_id")
        if oid:
            row = orders_model.get(db, oid)
        elif u.get("phone"):
            row = db.query_one(
                """SELECT o.id AS id, o.customer_id AS customer_id FROM orders o
                   JOIN customers c ON c.id = o.customer_id
                   WHERE c.phone = ? AND o.status IN
                   ('pending','confirmed','waiting_deposit','shipped')
                   ORDER BY o.id DESC LIMIT 1""", (u["phone"],))
        else:
            row = None
        if row is None:
            unmatched += 1
            continue
        orders_model.set_status(db, row["id"], status)
        trust.refresh(db, row["customer_id"])
        updated += 1
    return {"updated": updated, "unmatched": unmatched, "unmapped": unmapped}
